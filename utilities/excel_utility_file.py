import os
import time
import zipfile
import tempfile
import openpyxl
from contextlib import contextmanager
from openpyxl.styles import Font, PatternFill


REQUIRED_SHEETS = [
    "Login",
    "Register",
    "Logout",
    "Search",
    "Add to Cart",
    "Checkout",
    "Forgot Password",
]


@contextmanager
def _workbook_lock(file):
    lock_file = f"{file}.lock"
    os.makedirs(os.path.dirname(os.path.abspath(lock_file)), exist_ok=True)

    with open(lock_file, "a+b") as lock_handle:
        lock_handle.seek(0, os.SEEK_END)
        if lock_handle.tell() == 0:
            lock_handle.write(b"0")
            lock_handle.flush()

        if os.name == "nt":
            import msvcrt

            while True:
                try:
                    lock_handle.seek(0)
                    msvcrt.locking(lock_handle.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError:
                    time.sleep(0.05)
            try:
                yield
            finally:
                lock_handle.seek(0)
                msvcrt.locking(lock_handle.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl

            fcntl.flock(lock_handle.fileno(), fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(lock_handle.fileno(), fcntl.LOCK_UN)


def _save_workbook_atomically(workbook, file):
    directory = os.path.dirname(os.path.abspath(file))
    file_descriptor, temporary_file = tempfile.mkstemp(
        prefix=".workbook-", suffix=".xlsx", dir=directory
    )
    os.close(file_descriptor)
    try:
        workbook.save(temporary_file)
        os.replace(temporary_file, file)
    finally:
        if os.path.exists(temporary_file):
            os.remove(temporary_file)


def _ensure_workbook(file):
    os.makedirs(os.path.dirname(file) or ".", exist_ok=True)

    if not os.path.exists(file) or os.path.getsize(file) == 0:
        workbook = openpyxl.Workbook()
        if workbook.sheetnames and workbook.sheetnames[0] == "Sheet":
            workbook.remove(workbook["Sheet"])
        for sheet_name in REQUIRED_SHEETS:
            if sheet_name not in workbook.sheetnames:
                workbook.create_sheet(title=sheet_name)
        try:
            _save_workbook_atomically(workbook, file)
        finally:
            workbook.close()
        return

    try:
        with zipfile.ZipFile(file, "r") as archive:
            corrupt_member = archive.testzip()
            if corrupt_member is not None:
                raise zipfile.BadZipFile(
                    f"Corrupted workbook member: {corrupt_member}"
                )
    except (zipfile.BadZipFile, OSError, RuntimeError):
        workbook = openpyxl.Workbook()
        if workbook.sheetnames and workbook.sheetnames[0] == "Sheet":
            workbook.remove(workbook["Sheet"])
        for sheet_name in REQUIRED_SHEETS:
            if sheet_name not in workbook.sheetnames:
                workbook.create_sheet(title=sheet_name)
        try:
            _save_workbook_atomically(workbook, file)
        finally:
            workbook.close()
        return

    workbook = openpyxl.load_workbook(file)
    try:
        if any(sheet_name not in workbook.sheetnames for sheet_name in REQUIRED_SHEETS):
            for sheet_name in REQUIRED_SHEETS:
                if sheet_name not in workbook.sheetnames:
                    workbook.create_sheet(title=sheet_name)
            _save_workbook_atomically(workbook, file)
    finally:
        workbook.close()


def getrow_count(file, sheetname):
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        return sheet.max_row
    finally:
        workbook.close()


def getcol_count(file, sheetname):
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        return sheet.max_column
    finally:
        workbook.close()


def read_data(file, sheetname, rowno, colno):
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        return sheet.cell(rowno, colno).value
    finally:
        workbook.close()


def write_data(file, sheetname, rowno, colno, data):
    try:
        workbook = openpyxl.load_workbook(file)
        try:
            sheet = workbook[sheetname]
            sheet.cell(rowno, colno).value = data
            workbook.save(file)
        finally:
            workbook.close()
    except PermissionError:
        return


def fill_green(file, sheetname, rowno, colno):
    try:
        workbook = openpyxl.load_workbook(file)
        try:
            sheet = workbook[sheetname]
            green = PatternFill(start_color='60b212', end_color='60b212', fill_type='solid')
            sheet.cell(rowno, colno).fill = green
            sheet.cell(rowno, colno).font = Font(color='FFFFFF')
            workbook.save(file)
        finally:
            workbook.close()
    except PermissionError:
        return


def fill_red(file, sheetname, rowno, colno):
    try:
        workbook = openpyxl.load_workbook(file)
        try:
            sheet = workbook[sheetname]
            red = PatternFill(start_color='ff0000', end_color='ff0000', fill_type='solid')
            sheet.cell(rowno, colno).fill = red
            sheet.cell(rowno, colno).font = Font(color='FFFFFF')
            workbook.save(file)
        finally:
            workbook.close()
    except PermissionError:
        return


def update_case_result(file, sheetname, rowno, actual_result, passed):
    try:
        with _workbook_lock(file):
            _ensure_workbook(file)
            workbook = openpyxl.load_workbook(file)
            try:
                sheet = workbook[sheetname]
                sheet.cell(rowno, 8).value = actual_result
                result_cell = sheet.cell(rowno, 10)
                result_cell.value = "PASS" if passed else "FAIL"
                result_cell.fill = PatternFill(
                    start_color="60b212" if passed else "ff0000",
                    end_color="60b212" if passed else "ff0000",
                    fill_type="solid",
                )
                result_cell.font = Font(color="FFFFFF")
                _save_workbook_atomically(workbook, file)
            finally:
                workbook.close()
    except PermissionError:
        return


def get_data_from_excel_file(file, sheetname, minrownum,maxrownum):
    final_list = []
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        for row in sheet.iter_rows(min_row=minrownum, max_row=maxrownum, values_only=True):
            testcaseid,first_name,last_name,email,telephone,password,password_confirm,news_letter_options,privacy_policy_checkbox_text=row
            final_list.append((str(testcaseid or ""),str(first_name or ""),str(last_name or ""),
                               str(email or ""),str(telephone or ""),
                               str(password or ""),str(password_confirm or ""),
                               str(news_letter_options or ""), str(privacy_policy_checkbox_text or "")))
    finally:
        workbook.close()
    return final_list


def get_login_data_from_excel_file(file, sheetname, minrownum, maxrownum):
    final_list = []
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        for row in sheet.iter_rows(min_row=minrownum, max_row=maxrownum, values_only=True):
            testcaseid, email, password = row
            final_list.append(
                (str(testcaseid or ""), str(email or ""), str(password or ""))
            )
    finally:
        workbook.close()
    return final_list


def get_forgot_data_from_excel_file(file, sheetname, minrownum, maxrownum):
    final_list = []
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        for row in sheet.iter_rows(min_row=minrownum, max_row=maxrownum, values_only=True):
            testcaseid, email = row
            final_list.append((str(testcaseid or ""), str(email or "")))
    finally:
        workbook.close()
    return final_list


def get_search_data_from_excel_file(file, sheetname, minrownum, maxrownum):
    final_list = []
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        for row in sheet.iter_rows(min_row=minrownum, max_row=maxrownum, values_only=True):
            testcaseid, product_name, category, email, password = row
            final_list.append(
                (
                    str(testcaseid or ""),
                    str(product_name or ""),
                    str(category or ""),
                    str(email or ""),
                    str(password or ""),
                )
            )
    finally:
        workbook.close()
    return final_list


def get_add_to_cart_data_from_excel_file(file, sheetname, minrownum, maxrownum):
    final_list = []
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        for row in sheet.iter_rows(min_row=minrownum, max_row=maxrownum, values_only=True):
            testcaseid, product_name = row
            final_list.append((str(testcaseid or ""), str(product_name or "")))
    finally:
        workbook.close()
    return final_list


def get_checkout_data_from_excel_file(file, sheetname, minrownum, maxrownum):
    final_list = []
    workbook = openpyxl.load_workbook(file)
    try:
        sheet = workbook[sheetname]
        for row in sheet.iter_rows(min_row=minrownum, max_row=maxrownum, values_only=True):
            testcaseid, product_name = row
            final_list.append((str(testcaseid or ""), str(product_name or "")))
    finally:
        workbook.close()
    return final_list
